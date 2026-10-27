"""Build a Part II record and optionally place Section 7.5 in the supplied PDF.

Section 7.5 occupies the unused space below Table 7.1 on original PDF page 20.
Re-running starts from the backed-up original, so the former end appendix is
removed and the main report stays at 40 pages. Other original pages are preserved.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
from html import escape
import importlib.metadata
import json
from pathlib import Path
import re
import shutil
import sys

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE / ".local" / "python"))


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def latest_summary(prefix: str) -> tuple[dict, str]:
    paths = sorted((HERE / "results").glob(prefix + "*/summary.json"))
    if not paths:
        return {}, "not run"
    return json.loads(paths[-1].read_text(encoding="utf-8")), paths[-1].parent.name


def make_content() -> tuple[list, dict]:
    offline, offline_dir = latest_summary("offline_")
    local, local_dir = latest_summary("local_services_")
    readiness, readiness_dir = latest_summary("readiness_")
    state_file = HERE / "results" / "lab_status.json"
    lab = json.loads(state_file.read_text(encoding="utf-8-sig")) if state_file.exists() else {}
    now = dt.datetime.now(dt.timezone(dt.timedelta(hours=5, minutes=30))).isoformat(timespec="seconds")
    blocks = []

    def para(title, body):
        blocks.append((title, body))

    para("Scope and evidence status",
         "Prepared for Part II by Debapriya Kundu (CSE23115), " + now + ". "
         "This implementation record supplements Chapters 5, 6 and 7 and the lab foundation in "
         "Sections 11.2-11.4. It distinguishes proposed production controls, executed offline checks, "
         "local service checks and actual firewall-lab evidence. The report's expected-result "
         "cells are not observations. Joint experiments E1-E8 remain pending unless their real captures "
         "are explicitly cited. No hospital production network or patient data was used.")
    blocks.append(("Original source preservation", [
        ["Item", "Recorded status"],
        ["Existing tracked source files changed by this work", str(len(offline.get("existing_source_files_changed", []))) if offline else "not checked"],
        ["Offline evidence", "execution/results/" + offline_dir],
        ["Local service evidence", "execution/results/" + local_dir],
        ["Firewall lab evidence verified", str(bool(lab.get("firewall_lab_verified", False)))],
        ["Current lab stage", lab.get("stage", "Installer preparation; no VM network result recorded")],
    ]))
    rows = [["Existing check", "Observed result"]]
    for check in offline.get("checks", []):
        rows.append([check["name"].replace("_", " "), "PASS (exit 0)" if check["exit_code"] == 0
                     else "FAILED (exit " + str(check["exit_code"]) + ")"])
    blocks.append(("Offline validation actually executed", rows))
    para("Meaning of the offline results",
         "The existing Part II unit suite covers 9 tests; the full existing suite contains 42 tests. "
         "The policy verification command exercises 30 simulator probes, and the subset checker examines "
         "18 lab rules including documented broader and lab-only cases. A passing result establishes "
         "behaviour of the written model and tools; it does not establish that a deployed firewall "
         "enforces the policy. Full command lines, stdout, stderr, exit codes and SHA256 manifests "
         "are retained in the cited results directory.")
    rows = [["Local integration check", "Observed result"]]
    for check in local.get("checks", []):
        rows.append([check["name"].replace("_", " "), "PASS" if check["passed"] else "FAILED"])
    blocks.append(("Original stand-in services exercised over sockets", rows))
    para("Service-test boundaries",
         "The original Python programs were started as separate processes on 127.0.0.1 with temporary "
         "ports. Checks exercised HTTP health and portal responses, TCP listener reachability, a valid "
         "HL7 observation and matching ACK, rejection of a malformed message, and the original MLLP client. "
         "These connections did not cross OPNsense. The EMR uses clear-text HTTP on port 443 in the lab "
         "so IPS can read it; it is not evidence of production HTTPS. The MSSQL stand-in accepts TCP "
         "connections and does not implement SQL or TDS. The HL7 service is a minimal demonstrator.")
    if readiness.get("passed"):
        capture_bytes = (HERE / "results" / readiness_dir / "firewall.pcap").stat().st_size
        pump_record = json.loads((HERE / "results" / readiness_dir / "pump-vm.json").read_text())
        ntp = next(item["observed"] for item in pump_record["checks"] if item["name"] == "Internal NTP response")
        rows = [["Actual VM readiness check group", "Observed result"]]
        for role in readiness["roles"]:
            rows.append([role["role"], str(role["checks"]) + " checks; " + ("PASS" if role["passed"] else "FAILED")])
        rows.append(["Loaded policy / temporary setup rules", str(readiness["policy_rule_count"]) + " / " + str(readiness["temporary_setup_rule_count"])])
        rows.append(["Blocked connections corroborated in PF packet log", str(sum(item["passed"] for item in readiness["blocked_flow_log_correlations"])) + " of 7"])
        blocks.append(("Real firewall-lab readiness observations", rows))
        para("Evidence and installed environment",
             "The lab was built with VirtualBox 7.2.22, OPNsense 25.7 (FreeBSD 14.3-RELEASE-p1), "
             "and four Alpine 3.20.10 guests running kernel 6.6.142-0-virt. Windows is the MGMT bastion, "
             "not a sixth VM. All five VMs were restarted after configuration; the guests then booted "
             "from their disks with their planned addresses. The test run is retained in execution/results/"
             + readiness_dir + ". Per-host JSON and transcripts, the active 18-rule API inventory, "
             "decoded PF log and a " + str(round(capture_bytes / 1024, 1)) + " KiB packet capture record the observations. All " + str(readiness["check_count"])
             + " readiness checks passed. The temporary package server was stopped, setup rules removed "
             "and stale zone states cleared before these checks.")
        para("What the real checks establish",
             "The MGMT host reached all four SSH services and the live database listener. The pump "
             "used internal UDP/TCP DNS, NTP and HL7 with an AA acknowledgement, while EMR, database "
             "and public web access were blocked. Radiology reached the EMR; its broker alias sent "
             "HL7, while the reading workstation's HL7 and database connections were blocked. Vendor "
             "DHCP assigned 10.61.0.50, firewall DNS and public HTTP/HTTPS worked, and internal EMR/DB "
             "access was blocked. All seven denied connections matched actual block entries in the "
             "PF packet log. DNS recursion used source 192.168.10.53, and upstream NTP used .57, "
             "as shown by pass entries. The measured internal NTP response was stratum " + str(ntp["stratum"])
             + " with leap status " + str(ntp["leap"]) + "; this is not proof of the production stratum-2 or NTS design.")
    elif readiness:
        para("Real firewall-lab readiness observations", "The latest readiness run did not pass. "
             "See execution/results/" + readiness_dir + "; unresolved checks must be corrected before handoff.")
    blocks.append(("Lab foundation: fixed interface and address mapping", [
        ["VM / interface", "Network and role"],
        ["opnsense em0 / WAN", "VirtualBox NAT for upstream access"],
        ["opnsense em1 / LAN", "MGMT host-only; firewall 192.168.90.1, bastion 192.168.90.5"],
        ["his-vm / em2 OPT1", "192.168.10.0/24; DB .20, EMR .21, interface engine .40, DNS .53, NTP .57"],
        ["pump-vm / em3 OPT2", "192.168.30.10/24; gateway 192.168.30.1"],
        ["rad-vm / em4 OPT3", "192.168.40.100/24 and broker alias .12; gateway .1"],
        ["vendor-vm / em5 OPT4", "10.61.0.0/24; DHCP .50-.99; gateway 10.61.0.1"],
    ]))
    para("Architecture and protocol consistency",
         "The original objects.yaml and production addresses are preserved. The production design "
         "contains ten zones; the lab represents HIS, DEVICES, RADIOLOGY, VENDOR-WIFI and MGMT. "
         "Internal clinical/device/research DNS flows use H_DNS; guest/vendor production policy permits "
         "public DNS, and this lab uses the firewall resolver for vendor clients. Therefore the protocol "
         "register's phrase 'all zones to H_DNS only' needs this explicit exception when interpreting "
         "Chapter 6. DNSSEC, authenticated NTP/NTS, TLS logging and TLS interception are proposed "
         "production controls, not demonstrated by the local service tests.")
    para("CIA argument and controls outside the lab",
         "Confidentiality depends on scoped access and appropriate encryption; a permitted port alone "
         "does not supply encryption. Integrity depends on narrow permissions, protocol validation and "
         "the inspection engines actually enabled. Availability depends on correctly working DNS, time, "
         "state tracking and operational change control. HA, NAC/802.1X and private VLAN enforcement "
         "remain design/configuration references; this five-VM lab cannot demonstrate physical-switch "
         "peer isolation or HA failover. No claim of measured production performance is made.")
    para("Design assumptions requiring team review",
         "The report describes over 1,000 clinical PCs while Z_CLINICAL is a /23 with 510 usable "
         "addresses, including its gateway. The original brief does not supply that endpoint count. "
         "The capacity assumption must be clarified before a production rollout; this work does not "
         "change addressing behind Part III's rules. The brief's illustrative radiology-to-SQL permit "
         "is intentionally replaced by denial plus broker HL7 in the team design. Retain the Chapter 9 "
         "justification and validate the legitimate business-flow requirement with the policy owner.")
    para("Integration adaptations with original source preserved",
         "OPNsense 25.7 returns an empty list for missing alias UUIDs, rejects a comma-separated "
         "destination-port expression, and searches floating rules when no interface is supplied. "
         "Part II helpers pre-create the original aliases, represent exactly ports 80 and 443 with "
         "a port alias, and search the five original interfaces. Failed-load duplicates were removed; "
         "the final inventory contains exactly 18 intended rules. Windows CRLF was normalized only "
         "in deployed shell-script copies. Stopping DNS/NTP before initial configuration avoided an "
         "observed OpenRC restart race. Guest configuration binds outbound DNS/NTP to .53/.57. "
         "Interrupted package downloads were recovered using the same official signed APKs cached "
         "through Windows; Alpine signature verification remained enabled, repositories were restored "
         "and the temporary server stopped. Four temporary Windows routes send only lab-subnet "
         "traffic through 192.168.90.1; they do not replace the laptop's default route and must be "
         "re-created after a Windows reboot. All original tracked project files remain unchanged.")
    para("Sharing and reproduction",
         "All new project artefacts are under part2_architecture_protocols/execution. Downloaded media, "
         "temporary files and the unmodified input report backup are under ignored .local. Run "
         "run_offline_checks.py, fetch_lab_media.py and verify_local_services.py from this directory "
         "using Python, or follow execution/README.md from the repository root. Review the default "
         "dry run of create_lab_windows.ps1, then use -Create with an optional -VmDirectory. VM disks "
         "are local working data on a drive with enough space; their location is configurable and "
         "does not affect the addresses shared by the team.")
    para("Handoff and completion gate",
         "Part II hands over installed zone hosts and services, stable addressing, interface mapping, "
         "actual environment versions and evidence of readiness. Part III owns the firewall rule base. "
         "Part IV leads E1-E8 and the Observed cells in Table 11.2. A complete Part II VM handoff "
         "requires installed operating systems, working DNS/NTP/HTTP/MLLP, rules loaded and "
         "setup access removed, plus real permitted/blocked flow checks with firewall logs. "
         + ("Part II readiness has been demonstrated by the cited run. " if readiness.get("passed") else
            "Part II readiness remains pending. ") + "Shared E1-E8, their full captures, "
         "Part IV collector integration and the Observed cells remain the team's next work. App-ID, "
         "Suricata IPS, physical NAC/PVLAN and HA are not established by these readiness checks.")
    return blocks, {"recorded_at": now, "offline_results": offline_dir, "local_results": local_dir,
                    "lab_status": lab, "readiness_results": readiness_dir, "readiness": readiness}


def make_section_7_5(record: dict) -> dict:
    """Tabulate the recorded checks beside the Part II theory, without claiming E1-E8."""
    readiness = record["readiness"]
    if not readiness.get("passed"):
        raise RuntimeError("Section 7.5 requires a successful actual lab-readiness record")
    offline_dir = HERE / "results" / record["offline_results"]
    offline = json.loads((offline_dir / "summary.json").read_text(encoding="utf-8"))
    local = json.loads((HERE / "results" / record["local_results"] / "summary.json").read_text(encoding="utf-8"))
    checks = {check["name"]: check for check in offline["checks"]}
    required = {"part2_unit_tests", "all_existing_tests", "object_validation", "policy_lint",
                "policy_probes", "lab_subset"}
    if not required.issubset(checks) or any(checks[name]["exit_code"] for name in required):
        raise RuntimeError("Section 7.5 requires successful recorded software and policy checks")

    def observed_count(name: str, pattern: str) -> int:
        transcript = (offline_dir / checks[name]["transcript"]).read_text(encoding="utf-8")
        match = re.search(pattern, transcript)
        if not match:
            raise RuntimeError("No measured check count in " + name)
        return int(match.group(1))

    part2_tests = observed_count("part2_unit_tests", r"(\d+) passed")
    all_tests = observed_count("all_existing_tests", r"(\d+) passed")
    probes = observed_count("policy_probes", r"(\d+) probes, 0 failure")
    lab_rules = observed_count("lab_subset", r"(\d+) lab rules checked, 0 failure")
    local_passes = sum(check["passed"] for check in local["checks"])
    if not local["checks"] or local_passes != len(local["checks"]):
        raise RuntimeError("Section 7.5 requires successful recorded local-service checks")
    blocked = sum(item["passed"] for item in readiness["blocked_flow_log_correlations"])
    validation_rows = [
        ["Validation check", "Observed"],
        ["Unit tests: Part II / full suite", f"{part2_tests} / {all_tests} PASS"],
        ["Objects and policy lint", "PASS"],
        ["Simulator probes / lab rules", f"{probes} / {lab_rules} PASS"],
        ["Local service checks (loopback)", f"{local_passes} / {len(local['checks'])} PASS"],
        ["Policy / temporary setup rules", f"{readiness['policy_rule_count']} / {readiness['temporary_setup_rule_count']}"],
        ["Blocked flows matched in PF log", f"{blocked} / {len(readiness['blocked_flow_log_correlations'])} PASS"],
    ]
    labels = {"bastion": "MGMT bastion", "his-vm": "HIS VM", "pump-vm": "Device VM",
              "rad-vm": "Radiology VM", "vendor-vm": "Vendor VM"}
    readiness_rows = [["Lab role", "Checks", "Observed"]]
    for role in readiness["roles"]:
        readiness_rows.append([labels[role["role"]], str(role["checks"]), "PASS" if role["passed"] else "FAILED"])
    readiness_rows.append(["Total readiness", str(readiness["check_count"]), "PASS"])
    return {
        "title": "7.5 Implementation and readiness validation",
        "environment": "Implementation (11.2-11.4): VirtualBox 7.2.22, OPNsense 25.7, four Alpine 3.20.10 VMs; reboot verified.",
        "tables": [("Table 7.2: Software and policy checks", validation_rows),
                   ("Table 7.3: VM readiness checks", readiness_rows)],
        "notes": f"The {blocked} block checks are included in the {readiness['check_count']}. "
                 "E1-E8 remain separate. IPS/App-ID, NAC/PVLAN and HA are unverified. "
                 "EMR uses plaintext HTTP on 443; the database is a TCP stand-in.",
        "evidence": "part2_architecture_protocols/execution/results/ (offline, local_services and readiness records).",
    }


def make_section_overlay(path: Path, record: dict) -> str:
    """Use only the blank area under Table 7.1; refuse text that would reach the footer."""
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Paragraph, Table, TableStyle

    section = make_section_7_5(record)
    page_width, page_height = 595.276, 841.89
    left, width = 56.693, 481.89
    body_style = ParagraphStyle("SectionBody", fontName="Times-Roman", fontSize=9.5,
                                leading=11.2)
    evidence_style = ParagraphStyle("SectionEvidence", fontName="Times-Roman", fontSize=8.5,
                                    leading=10)
    items = []
    top = 196.0
    paragraph = Paragraph(escape(section["environment"]), body_style)
    _, height = paragraph.wrap(width, page_height)
    top -= height
    items.append((paragraph, left, top))
    caption_y = top - 8
    table_top = caption_y - 5
    captions = []
    table_bottoms = []
    for index, ((caption, rows), col_widths) in enumerate(zip(section["tables"], ([155, 80], [110, 42, 82.89]))):
        cells = [[Paragraph("<b>" + escape(cell) + "</b>" if row_index == 0 or (index == 1 and row_index == len(rows) - 1)
                            else escape(cell), body_style)
                  for cell in row] for row_index, row in enumerate(rows)]
        table = Table(cells, colWidths=col_widths)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6eef5")),
            ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#a8b2bc")),
            ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c7cdd3")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 1),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
        ]))
        table_width, height = table.wrap(sum(col_widths), page_height)
        x = left if index == 0 else left + 247
        items.append((table, x, table_top - height))
        captions.append((caption, x + table_width / 2))
        table_bottoms.append(table_top - height)
    top = min(table_bottoms) - 4
    paragraph = Paragraph(escape(section["notes"]), evidence_style)
    _, height = paragraph.wrap(width, page_height)
    top -= height
    items.append((paragraph, left, top))
    top -= 2
    paragraph = Paragraph("<b>Evidence:</b> " + escape(section["evidence"]), evidence_style)
    _, height = paragraph.wrap(width, page_height)
    top -= height
    items.append((paragraph, left, top))
    if top < 42:
        raise RuntimeError("Section 7.5 does not fit safely above the original page footer")
    drawing = canvas.Canvas(str(path), pagesize=(page_width, page_height))
    drawing.setFont("Times-Bold", 11.9552)
    drawing.drawString(left, 209.0, section["title"])
    drawing.setFont("Times-Roman", 9.5)
    for caption, centre in captions:
        drawing.drawCentredString(centre, caption_y, caption)
    for item, x, bottom in items:
        item.drawOn(drawing, x, bottom)
    drawing.save()
    markdown = ["# " + section["title"], "", section["environment"], ""]
    for caption, rows in section["tables"]:
        markdown.extend(["**" + caption + "**", "", "| " + " | ".join(rows[0]) + " |",
                         "|" + "---|" * len(rows[0])])
        markdown.extend("| " + " | ".join(row) + " |" for row in rows[1:])
        markdown.append("")
    markdown.extend([section["notes"], "", "Evidence: " + section["evidence"], ""])
    return "\n".join(markdown)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--update-main-report", action="store_true")
    args = parser.parse_args()
    from pypdf import PdfReader, PdfWriter
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
    from docx import Document
    from docx.shared import Inches, Pt

    report_dir = HERE / "report"
    report_dir.mkdir(exist_ok=True)
    blocks, record = make_content()
    title = "Part II - Implementation and Validation Record"
    markdown = ["# " + title, ""]
    styles = getSampleStyleSheet()
    styles["BodyText"].fontSize = 9
    styles["BodyText"].leading = 12
    styles["BodyText"].alignment = TA_LEFT
    story = [Paragraph(title, styles["Title"]), Spacer(1, 10)]
    word = Document()
    word.add_heading(title, 0)
    word.styles["Normal"].font.name = "Calibri"
    word.styles["Normal"].font.size = Pt(10)
    for heading, body in blocks:
        markdown += ["## " + heading, ""]
        story.append(Paragraph(escape(heading), styles["Heading2"]))
        word.add_heading(heading, 1)
        if isinstance(body, str):
            markdown += [body, ""]
            story += [Paragraph(escape(body), styles["BodyText"]), Spacer(1, 5)]
            word.add_paragraph(body)
        else:
            markdown += ["| " + " | ".join(body[0]) + " |", "|---|---|"]
            markdown += ["| " + " | ".join(row) + " |" for row in body[1:]] + [""]
            table = Table([[Paragraph(escape(str(c)), styles["BodyText"]) for c in row]
                           for row in body], colWidths=[210, 285], repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6eef5")),
                                       ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#b7c4cf")),
                                       ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                       ("LEFTPADDING", (0, 0), (-1, -1), 6),
                                       ("RIGHTPADDING", (0, 0), (-1, -1), 6)]))
            story += [table, Spacer(1, 8)]
            t = word.add_table(rows=0, cols=2)
            t.style = "Table Grid"
            for row in body:
                for cell, value in zip(t.add_row().cells, row):
                    cell.text = str(value)
    (report_dir / "part2_implementation_record.md").write_text("\n".join(markdown), encoding="utf-8")
    word.save(report_dir / "part2_implementation_record.docx")
    pdf = report_dir / "part2_implementation_record.pdf"

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(48, 28, "CS07 | Debapriya Kundu | Part II validation record")
        canvas.drawRightString(A4[0] - 48, 28, "P2-" + str(doc.page))
        canvas.restoreState()

    SimpleDocTemplate(str(pdf), pagesize=A4, rightMargin=48, leftMargin=48,
                      topMargin=42, bottomMargin=45).build(story, onFirstPage=footer, onLaterPages=footer)
    record["pdf_pages"] = len(PdfReader(pdf).pages)
    record["report_libraries"] = {name: importlib.metadata.version(name)
                                  for name in ("pypdf", "reportlab", "python-docx")}
    overlay = HERE / ".local" / "section_7_5_overlay.pdf"
    overlay.parent.mkdir(exist_ok=True)
    section_markdown = make_section_overlay(overlay, record)
    (report_dir / "section_7_5.md").write_text(section_markdown, encoding="utf-8")
    manifest_file = report_dir / "report_update.json"
    if args.update_main_report:
        main_pdf = ROOT / "SecureNetworks_CaseStudy_Report.pdf"
        baseline = json.loads((HERE / "baseline_sha256.json").read_text(encoding="utf-8"))
        original_sha = baseline["input_report_sha256"]
        last = json.loads(manifest_file.read_text(encoding="utf-8")) if manifest_file.exists() else {}
        if sha256(main_pdf) not in (original_sha, last.get("updated_report_sha256")):
            raise RuntimeError("The main PDF was changed separately; refusing to overwrite it")
        backup = HERE / ".local" / "original_report.pdf"
        backup.parent.mkdir(exist_ok=True)
        if not backup.exists():
            shutil.copyfile(main_pdf, backup)
        if sha256(backup) != original_sha:
            raise RuntimeError("Original report backup checksum mismatch")
        original = PdfReader(backup)
        page_index = 19
        if len(original.pages) != 40 or "7.4 CIA by zone" not in original.pages[page_index].extract_text():
            raise RuntimeError("The original report does not match the reviewed Section 7.5 placement")
        writer = PdfWriter(clone_from=original)
        writer.pages[page_index].merge_page(PdfReader(overlay).pages[0])
        merged = HERE / ".local" / "report_with_part2_record.pdf"
        writer.write(merged)
        reader = PdfReader(merged)
        assert len(reader.pages) == len(original.pages)
        for n, page in enumerate(original.pages):
            if n == page_index:
                assert reader.pages[n].extract_text().startswith(page.extract_text()), n
                assert "7.5 Implementation and readiness validation" in reader.pages[n].extract_text()
            else:
                assert page.get_contents().get_data() == reader.pages[n].get_contents().get_data(), n
        assert reader.page_labels == original.page_labels
        assert len(reader.outline) == len(original.outline)
        record.update({"original_report_sha256": original_sha,
                       "original_pages_preserved": len(original.pages) - 1,
                       "original_text_preserved_on_all_pages": True,
                       "modified_original_pdf_pages": [page_index + 1],
                       "section": "7.5 Implementation and readiness validation",
                       "section_pdf_page": page_index + 1,
                       "section_printed_page": "17",
                       "section_tables": ["7.2 Software and policy checks", "7.3 VM readiness checks"],
                       "placement": "Below Table 7.1 within Part II; no end appendix",
                       "updated_report_sha256": sha256(merged), "total_pages": len(reader.pages)})
        merged.replace(main_pdf)
        print("Updated Section 7.5 on PDF page 20 (printed page 17); main report has 40 pages")
        print("Removed the former end appendix; other 39 original page content streams are unchanged")
    manifest_file.write_text(json.dumps(record, indent=2) + "\n", encoding="utf-8")
    print("Report artefacts: " + str(report_dir))


if __name__ == "__main__":
    main()
